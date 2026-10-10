from __future__ import annotations

import csv
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BASE_DIR = Path(__file__).resolve().parents[1]
OUTPUT = BASE_DIR / "docs" / "May14_Loss_Analysis_and_Bot_Changes.docx"
TRADE_HISTORY = BASE_DIR / "logs" / "trade_history.csv"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: object, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run("" if text is None else str(text))
    run.bold = bold
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[object]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)
        set_cell_shading(table.rows[0].cells[index], "E8EEF5")

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)

    document.add_paragraph()


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def read_exit_rows() -> list[dict]:
    rows: list[dict] = []
    with TRADE_HISTORY.open(newline="") as csv_file:
        for row in csv.DictReader(csv_file):
            if row.get("event") == "EXIT" and row.get("date"):
                rows.append(row)
    return rows


def summarize(rows: list[dict]) -> dict:
    total = len(rows)
    pnl = round(sum(float(row.get("pnl") or 0) for row in rows), 2)
    wins = sum(float(row.get("pnl") or 0) > 0 for row in rows)
    losses = sum(float(row.get("pnl") or 0) < 0 for row in rows)
    return {
        "trades": total,
        "pnl": pnl,
        "wins": wins,
        "losses": losses,
        "win_rate": round((wins / total) * 100, 2) if total else 0,
        "target_percent": Counter(row.get("target_percent") for row in rows),
        "exit_profile": Counter(row.get("exit_profile") for row in rows),
        "news": Counter((row.get("news_sentiment"), row.get("news_risk_level")) for row in rows),
        "underlying": Counter(row.get("underlying") or row.get("option_symbol", "")[:6] for row in rows),
        "reason": Counter(row.get("reason") for row in rows),
    }


def daily_summary(rows: list[dict]) -> list[list[object]]:
    by_day = defaultdict(lambda: {"trades": 0, "pnl": 0.0, "wins": 0, "losses": 0})
    for row in rows:
        pnl = float(row.get("pnl") or 0)
        day = row["date"]
        by_day[day]["trades"] += 1
        by_day[day]["pnl"] += pnl
        by_day[day]["wins"] += pnl > 0
        by_day[day]["losses"] += pnl < 0

    output = []
    for day in sorted(by_day):
        item = by_day[day]
        output.append([
            day,
            item["trades"],
            item["wins"],
            item["losses"],
            round(item["pnl"], 2),
        ])
    return output


def build_document() -> None:
    rows = read_exit_rows()
    before = [row for row in rows if row["date"] <= "2026-05-14"]
    after = [row for row in rows if row["date"] > "2026-05-14"]
    before_summary = summarize(before)
    after_summary = summarize(after)

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Calibri"
    styles["Title"].font.size = Pt(20)
    styles["Heading 1"].font.name = "Calibri"
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 2"].font.name = "Calibri"
    styles["Heading 2"].font.size = Pt(12.5)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Trading Bot Loss Review After 14 May").bold = True
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Evidence, likely causes, and changes made for safer paper trading").italic = True

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(
        "The trading logs show that performance materially weakened after 14 May 2026. "
        "The main issue was not only lower win rate; the larger issue was that losing trades became much larger. "
        "After 14 May, the bot mostly traded setups labeled STRONG_MOMENTUM, used wider 18% stop losses, "
        "and earlier versions allowed high-risk news, re-entry after stop loss, and oversized single trades."
    )

    add_table(
        document,
        ["Period", "Closed trades", "Wins", "Losses", "Win rate", "Total P&L"],
        [
            ["On/before 14 May", before_summary["trades"], before_summary["wins"], before_summary["losses"], f"{before_summary['win_rate']}%", f"Rs {before_summary['pnl']}"],
            ["After 14 May", after_summary["trades"], after_summary["wins"], after_summary["losses"], f"{after_summary['win_rate']}%", f"Rs {after_summary['pnl']}"],
        ],
    )

    document.add_heading("Daily P&L Snapshot", level=1)
    add_table(document, ["Date", "Closed trades", "Wins", "Losses", "P&L"], daily_summary(rows))

    document.add_heading("Likely Reasons Losses Increased After 14 May", level=1)
    add_bullets(
        document,
        [
            "Loss size increased sharply. Several post-14 May losses were larger than Rs 10,000, while many earlier losses were much smaller.",
            "The bot over-trusted STRONG_MOMENTUM. All post-14 May closed trades were marked STRONG_MOMENTUM, but strong ADX/setup score can also appear after a move is already stretched.",
            "Older target settings were too ambitious. Four post-14 May trades used a 60% target with an 18% stop loss, so trades often needed a large continuation move to win.",
            "High-risk news was previously not blocking trades. May 19 and some other trades were taken even when news risk was HIGH_RISK/HIGH.",
            "The bot allowed same-direction re-entry after SL. On May 20 it re-entered SENSEX PE after a SENSEX PE stop loss, increasing drawdown.",
            "Position sizing allowed one trade to carry large exposure. This amplified the impact when the market reversed.",
            "Market regime appears to have become more choppy/reversal-heavy after 14 May. Momentum entries that worked earlier started reversing before final targets.",
        ],
    )

    document.add_heading("Changes Made", level=1)
    add_table(
        document,
        ["Area", "Change", "Why it helps"],
        [
            ["Targets", "Added staged targets: T1, T2, T3.", "Protects gains earlier instead of waiting for one final target."],
            ["SL movement", "T1 moves SL to entry; T2 moves SL to T1; T3 exits full trade.", "Turns profitable movement into protected trades."],
            ["News filter", "Enabled news blocking and high-risk news blocking.", "Avoids trading when headlines conflict or risk is high."],
            ["Cooldown", "Added 90-minute same underlying/same direction cooldown after negative SL.", "Prevents repeated revenge-style re-entry."],
            ["Risk mode", "After daily risk breach, bot continues analysis but logs signals as BLOCKED_DAILY_RISK.", "Allows learning without adding losses."],
            ["Paper capital", "Paper balance now rolls forward from starting Rs 100,000 plus realized P&L.", "Next-day sizing reflects actual paper account performance."],
            ["Single-index trades", "Allowed only when strong enough and at reduced 50% size.", "Captures opportunity without removing confirmation safety entirely."],
            ["Open-trade monitor", "Open trade target/SL checks run faster than fresh analysis.", "Improves SL/target responsiveness."],
        ],
    )

    document.add_heading("Current Exit Plan", level=1)
    add_table(
        document,
        ["Profile", "Target 1", "Target 2", "Target 3 / Exit", "Initial SL"],
        [
            ["Strong momentum", "15%", "25%", "35%", "18%"],
            ["Normal momentum", "10%", "18%", "25%", "18%"],
            ["Weak/risky momentum", "8%", "12%", "18%", "12%"],
        ],
    )

    document.add_heading("Current Risk Controls", level=1)
    add_bullets(
        document,
        [
            "Total open fund usage can go up to 70% of rolling paper balance.",
            "Single trade can also use up to 70%, as requested, while risk limit still controls allowed lots.",
            "Daily failed-trade limit remains active.",
            "Same-index and same-direction negative SL triggers a 90-minute cooldown.",
            "High-risk news and news conflict can block trades.",
            "When daily risk is hit, the bot keeps producing trade ideas but marks them analysis-only.",
        ],
    )

    document.add_heading("Recommended Watch Points", level=1)
    add_bullets(
        document,
        [
            "Track whether STRONG_MOMENTUM still loses often. If yes, raise strong thresholds or add an overextension filter.",
            "Watch entries after big gap-up or gap-down opens; these may reverse despite high ADX.",
            "Review whether T1/T2 protect enough profit over the next 5 to 10 trading days.",
            "Before live auto, confirm broker-side or stronger live SL handling.",
            "Keep paper testing until the rolling paper account curve stabilizes.",
        ],
    )

    document.add_heading("Bottom Line", level=1)
    document.add_paragraph(
        "After 14 May, the bot appears to have faced a market regime where strong-looking momentum setups reversed more often. "
        "The code was also too permissive: large single-trade exposure, high-risk news allowed, no same-direction SL cooldown, "
        "and targets that were too far. The new changes are designed to make the bot more defensive, protect profit earlier, "
        "and keep analysis running even when trading is blocked."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
