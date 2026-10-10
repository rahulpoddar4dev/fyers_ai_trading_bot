from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parents[1]
OUTPUT = BASE_DIR / "Trading_Bot_Project_Documentation.docx"


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
INK = "18212F"
MUTED = "667085"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_pr = table._tbl.tblPr
    tbl_w = table_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = table_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        table_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_text(run, bold=False, color=INK, size=None):
    run.font.name = "Calibri"
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AI Trading Bot Project Documentation")
    set_text(r, bold=True, color=BLUE, size=22)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Fyers API + Python + Paper Trading Dashboard")
    set_text(r, color=DARK_BLUE, size=13)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Prepared on {date.today().isoformat()}")
    set_text(r, color=MUTED, size=10)


def add_note(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    set_text(r, bold=True, color=DARK_BLUE)
    r = p.add_run(body)
    set_text(r)
    doc.add_paragraph()


def add_kv_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2700, 6660])
    header = table.rows[0].cells
    header[0].text = "Item"
    header[1].text = "Details"
    for cell in header:
        set_cell_shading(cell, LIGHT_BLUE)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_text(run, bold=True, color=DARK_BLUE)

    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
        for cell in cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    set_text(run)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def build_doc():
    doc = Document()
    configure_styles(doc)
    add_title(doc)

    add_note(
        doc,
        "Safety note",
        "This project is currently configured for paper trading. It logs decisions and simulated trades. "
        "Live auto trading should remain disabled until backtesting, paper trading, and broker-side risk checks are reviewed.",
    )

    doc.add_heading("1. Project Overview", level=1)
    doc.add_paragraph(
        "We built a local Python trading bot that connects to Fyers API v3, analyzes NIFTY and SENSEX, "
        "selects option contracts, tracks paper trades, records results, and displays activity in a browser dashboard."
    )
    add_kv_table(doc, [
        ("Project folder", "trading-bot"),
        ("Primary command", "python3 main.py"),
        ("Dashboard", "Started automatically by main.py on a local 127.0.0.1 port"),
        ("Current mode", "PAPER mode; no real orders are placed while PAPER_TRADE is True"),
        ("AI mode", "Mock Claude mode is enabled until Anthropic API credits are available"),
    ])

    doc.add_heading("2. Major Milestones Completed", level=1)
    add_numbered(doc, [
        "Created the Python trading-bot project structure.",
        "Installed and wired Fyers API v3, Anthropic, pandas, and request-related dependencies.",
        "Built Fyers login, saved-session reuse, and browser-based auth_code capture fallback.",
        "Added live quote and historical candle fetching for NIFTY and SENSEX.",
        "Added rule-based strategy logic using EMA, RSI, VWAP, ATR, and ADX.",
        "Added NIFTY versus SENSEX comparison before selecting a trade.",
        "Added option selection for NFO/BFO option contracts.",
        "Added paper trading, CSV logging, date-wise logs, and open-trade tracking.",
        "Added a browser dashboard for current fund, account, open trade, today trades, history, and performance.",
        "Added risk management, trailing stop loss, adaptive target/SL, and dynamic paper balance testing.",
        "Added a backtesting script for strategy validation against historical candles.",
    ])

    doc.add_heading("3. Important Files", level=1)
    add_kv_table(doc, [
        ("config.py", "Main settings: mode, paper balance, risk rules, Fyers redirect URI, news settings, indicators, timing rules."),
        ("main.py", "Starts dashboard, logs in to Fyers, runs analysis cycle, applies risk checks, and executes paper/live flow."),
        ("fyers_login.py", "Handles saved sessions, automatic callback login, redirect URL fallback, and session token generation."),
        ("market_data.py", "Fetches quote and candle data for NIFTY/SENSEX."),
        ("strategy.py", "Builds BUY/SELL/HOLD signals using indicators and filters."),
        ("option_selector.py", "Chooses option contract, strike, expiry, target, stop loss, and adaptive exit profile."),
        ("position_sizer.py", "Calculates lots/quantity from fund usage and dynamic risk limits."),
        ("risk_manager.py", "Blocks new trades based on daily stop losses, loss limit, profit target, and trade guardrails."),
        ("trade_state.py", "Tracks one open trade, monitors target/SL, square-off, stale contracts, and trailing SL."),
        ("dashboard.py / dashboard.html", "Local browser dashboard server and UI."),
        ("paper_trader.py / trade_history.py", "CSV logging for decisions, entries, exits, P&L, and day-wise records."),
        ("backtest_strategy.py", "Historical candle backtest runner."),
    ])

    doc.add_heading("4. Current Trading Flow", level=1)
    add_numbered(doc, [
        "Start main.py and dashboard.",
        "Reuse saved Fyers session if still valid; otherwise complete Fyers login.",
        "Check account balance and use paper override if PAPER mode is active.",
        "Block fresh trades during unsafe time windows such as before 09:25, lunch chop, after 15:00, expiry-day late window, or square-off time.",
        "Fetch NIFTY and SENSEX quotes/candles.",
        "Calculate EMA9, EMA21, RSI14, VWAP, ATR14, and ADX14.",
        "Build strategy signal for each index.",
        "Compare NIFTY and SENSEX; take trade only when confirmation rules are satisfied.",
        "Fetch news sentiment for analysis only; news does not block trades now.",
        "Select option contract and calculate adaptive target/SL.",
        "Size position using 70% fund usage and dynamic risk rules.",
        "Record paper trade and track open trade until target, stop loss, trailing SL, or square-off.",
    ])

    doc.add_heading("5. Strategy Logic", level=1)
    add_kv_table(doc, [
        ("BUY setup", "Close above EMA9, EMA21, and VWAP with RSI in bullish range; trend and volatility filters must pass."),
        ("SELL setup", "Close below EMA9, EMA21, and VWAP with RSI in bearish range; trend and volatility filters must pass."),
        ("Sideways filter", "Avoids trades when EMA gap is too small."),
        ("Low volatility filter", "Avoids trades when ATR percentage is below configured minimum."),
        ("Trend filter", "Avoids trades when ADX is below configured minimum."),
        ("Score filter", "Rejects weak setups below MIN_SETUP_SCORE."),
        ("NIFTY/SENSEX confirmation", "Bot compares both indices; if only one index signals or signals conflict, it can hold depending on config."),
    ])

    doc.add_heading("6. Risk Management", level=1)
    add_kv_table(doc, [
        ("Paper balance", "PAPER_BALANCE_OVERRIDE = 100000 for testing lots without real fund impact."),
        ("Fund usage", "MAX_FUND_USAGE_PERCENT = 70. Bot can use up to 70% of available/paper fund for one trade."),
        ("Fixed trade cap", "MAX_TRADE_AMOUNT = 0, meaning no fixed cap; fund percentage controls premium usage."),
        ("Per-trade risk", "Dynamic: MAX_RISK_PER_TRADE_PERCENT_OF_USED_FUND = 18% of used fund."),
        ("Daily loss limit", "Dynamic: MAX_DAILY_LOSS_PERCENT_OF_USED_FUND = 25% of used fund."),
        ("Daily profit target", "Dynamic: MAX_DAILY_PROFIT_PERCENT_OF_USED_FUND = 50% of used fund."),
        ("Stop-loss guard", "MAX_DAILY_STOP_LOSSES = 2. After 2 stop losses in one day, fresh trades stop."),
        ("Open trade limit", "MAX_OPEN_TRADES = 1. Bot tracks one open trade at a time while testing."),
    ])

    doc.add_heading("7. Adaptive Target and Stop Loss", level=1)
    add_kv_table(doc, [
        ("Strong momentum", "Target 60%, stop loss 18% when setup score and ADX are strong."),
        ("Normal momentum", "Target 35%, stop loss 18%."),
        ("Weak/risky setup", "Target 24%, stop loss 12%."),
        ("Trailing SL", "Enabled. Profit 15% moves SL to entry; profit 25% enables trailing stop with 12% distance."),
        ("Dashboard fields", "exit_profile, target_percent, stop_loss_percent, and exit profile reason are logged/displayed."),
    ])

    doc.add_heading("8. Option Selection", level=1)
    add_bullets(doc, [
        "BUY signal selects a CE option.",
        "SELL signal selects a PE option.",
        "NIFTY options are selected from NFO.",
        "SENSEX options are selected from BFO.",
        "Readable option names are shown, for example: NIFTY 19 MAY 23750 PE.",
        "OPTION_SELECTION_MODE = AUTO currently prefers safer ITM-style selection instead of pure ATM.",
    ])

    doc.add_heading("9. Login and Dashboard", level=1)
    add_kv_table(doc, [
        ("Saved session", "fyers_session.json stores same-day access token. Bot should reuse it until token expires."),
        ("Automatic login", "Local callback server or redirect URI captures auth_code."),
        ("Chrome fallback", "If Fyers redirects to redirect URI, code can extract auth_code from the URL."),
        ("Dashboard account display", "Dashboard shows the Fyers account name/user id from profile."),
        ("Dashboard run mode", "main.py starts the dashboard automatically; dashboard.py separately is normally not needed."),
    ])

    doc.add_heading("10. News Analysis", level=1)
    add_bullets(doc, [
        "News RSS headlines are fetched for analysis/logging.",
        "News sentiment and risk level are saved in CSV and shown in dashboard.",
        "News does not currently block chart-based trades because NEWS_BLOCK_TRADES = False.",
        "SSL fallback was added for macOS certificate issues.",
    ])

    doc.add_heading("11. Logs and Reporting", level=1)
    add_kv_table(doc, [
        ("Paper decisions", "logs/paper_option_trades.csv and logs/by_date/YYYY-MM-DD/paper_option_trades.csv"),
        ("Trade history", "logs/trade_history.csv and logs/by_date/YYYY-MM-DD/trade_history.csv"),
        ("Open trade state", "open_trade.json"),
        ("Account status", "account_status.json"),
        ("Performance", "performance_report.py summarizes closed trades, wins, losses, win rate, and P&L."),
    ])

    doc.add_heading("12. Testing Commands", level=1)
    add_kv_table(doc, [
        ("Run bot", "python3 main.py"),
        ("Run dashboard only", "python3 dashboard.py"),
        ("Run performance report", "python3 performance_report.py"),
        ("Run backtest", "python3 backtest_strategy.py"),
        ("Force fresh login", "rm fyers_session.json, then python3 main.py"),
    ])

    doc.add_heading("13. Current Safety Status", level=1)
    add_note(
        doc,
        "Current status",
        "The project is suitable for paper trading and observation. Live trading should remain disabled until backtesting, "
        "paper logs, option fills, slippage, brokerage, and SEBI/static-IP requirements are reviewed.",
    )
    add_bullets(doc, [
        "PAPER_TRADE is expected to remain True while testing.",
        "ENABLE_LIVE_AUTO should remain False until explicitly ready.",
        "Do not commit or share API secrets.",
        "Move keys from config.py to environment variables before serious use.",
        "Validate SENSEX/BFO availability with Fyers account permissions.",
    ])

    doc.add_heading("14. Recommended Next Improvements", level=1)
    add_bullets(doc, [
        "Run at least 2-4 weeks of paper trading logs before live mode.",
        "Compare backtest results with live paper results.",
        "Add brokerage/slippage estimates into P&L.",
        "Add dashboard cards for adaptive target profile and daily dynamic risk limits.",
        "Move API keys to .env and remove hardcoded secrets.",
        "Add unit tests for risk sizing, target/SL profiles, and market comparator decisions.",
    ])

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Trading Bot Project Documentation")
    set_text(run, color=MUTED, size=9)

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_doc())
